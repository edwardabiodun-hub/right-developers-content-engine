"""Create review-ready Word files from a generated social content package."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


PLATFORM_LABELS = {
    "linkedin": "LinkedIn",
    "instagram": "Instagram",
    "facebook": "Facebook",
}

VISUAL_CAPTIONS = {
    "linkedin": "Specification before terms.",
    "instagram": "A serious inquiry starts with detail.",
    "facebook": "A price request is not yet a commercial inquiry.",
}


def _set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    # Kept here for future table-based campaign summaries; the current layout is intentionally simple.
    _ = (cell, top, start, bottom, end)


def _configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(36, 44, 46)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    if "Small Meta" not in [style.name for style in styles]:
        meta = styles.add_style("Small Meta", WD_STYLE_TYPE.PARAGRAPH)
        meta.font.name = "Aptos"
        meta.font.size = Pt(9)
        meta.font.color.rgb = RGBColor(93, 105, 106)
    else:
        meta = styles["Small Meta"]
    meta.paragraph_format.space_after = Pt(4)


def _add_text_line(document: Document, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        return
    if stripped.startswith("- "):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(stripped[2:])
        return
    if stripped.startswith("Slide ") or stripped == "Caption":
        document.add_heading(stripped, level=2)
        return
    document.add_paragraph(stripped)


def _add_post_copy(document: Document, source: Path) -> None:
    document.add_heading("Suggested Post Copy", level=1)
    for line in source.read_text(encoding="utf-8").splitlines():
        _add_text_line(document, line)


def create_platform_document(platform: str, package_dir: Path, output_dir: Path) -> Path:
    label = PLATFORM_LABELS[platform]
    document = Document()
    _configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.add_paragraph(f"{label} Trade Readiness Post", style="Title")
    meta = document.add_paragraph(style="Small Meta")
    meta.add_run("Right Developers and Investment Group Inc. | Campaign TR-001 | Human review required")

    image = package_dir / "assets" / f"{platform}.png"
    if not image.exists():
        raise FileNotFoundError(f"Missing platform image: {image}")

    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(image), width=Inches(5.9 if platform == "instagram" else 6.35))

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(14)
    run = caption.add_run(VISUAL_CAPTIONS[platform])
    run.italic = True
    run.font.color.rgb = RGBColor(86, 96, 95)

    document.add_heading("Visual Caption", level=1)
    document.add_paragraph(
        f'Use the embedded {label} image with this caption: “{VISUAL_CAPTIONS[platform]}”'
    )
    _add_post_copy(document, package_dir / f"{platform}.md")

    output_dir.mkdir(parents=True, exist_ok=True)
    destination = output_dir / f"{platform}-trade-readiness-post.docx"
    document.save(destination)
    return destination


def main() -> int:
    parser = argparse.ArgumentParser(description="Create Word documents with embedded social post visuals.")
    parser.add_argument("--package-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    for platform in PLATFORM_LABELS:
        create_platform_document(platform, args.package_dir, args.output_dir)
    print(f"Word package written to {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
